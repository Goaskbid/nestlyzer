"""Builds the release report and deploy guide as .docx and .pdf into dist/ (pandoc + LibreOffice)."""
import json, os, re, shutil, subprocess, sys
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
W = os.path.join(ROOT, 'dist', '_docs'); os.makedirs(W, exist_ok=True)
shots = os.path.join(ROOT, 'tests', 'shots')
for f in os.listdir(shots):
    if f.startswith(('final_', 'sa_', 'v21_', 'v22_', 'v23_', 'v24_')): shutil.copy(os.path.join(shots, f), W)
for f in ('old_desktop_top.png', 'old_mobile_top.png'):
    src = os.path.join('/home/claude/audit/shots', f)
    if os.path.exists(src): shutil.copy(src, W)

SRC = os.environ.get('NZ_REPORT', 'RELEASE_2.4.md')
VER = re.search(r'(\d+\.\d+)', SRC).group(1)
md = open(os.path.join(ROOT, 'docs', SRC)).read()
d = json.load(open(os.path.join(ROOT, 'tests', 'results', 'regression.json')))
rows = ['| Check | What is tested | Result | Evidence |', '|---|---|---|---|']
for r in d['results']:
    st = {'PASS': 'Pass', 'FAIL': 'FAIL'}.get(r['status'], 'Not verifiable here')
    ev = re.sub(r'\s+', ' ', r['evidence']).replace('|', '/')
    ev = ev[:157] + '...' if len(ev) > 160 else ev
    rows.append('| %s | %s | %s | %s |' % (r['id'], re.sub(r'\s+', ' ', r['name']).replace('|', '/'), st, ev))
table = 'Result: **%d pass, %d fail, %d not verifiable offline** (%s).\n\n' % (d['pass'], d['fail'], d['other'], d['when']) + '\n'.join(rows)
md = md.replace('REGRESSION_TABLE', 'Run in headless Chromium on desktop and phone sizes against the built site.\n\n' + table)
at = os.path.join(ROOT, 'docs', '_audit_table.md')
if os.path.exists(at): md = md.replace('AUDIT_TABLE', open(at).read())
sa = os.path.join(ROOT, 'tests', 'results', 'standalone.json')
if os.path.exists(sa):
    j = json.load(open(sa))
    md = md.replace('STANDALONE', '`nestlyzer.html` was opened straight from disk in Chromium at desktop (1440 x 900) and phone (390 x 844) sizes. Results: ' + '; '.join('%s: %s' % (k.replace('_', ' '), v) for k, v in j.items()) + '.')

shots_md = '''## Home page: old and new

![Old v86](old_desktop_top.png){width=100%}

![New 2.1](v21_d_home.png){width=100%}

## Explore with a place open: why-card and fact tiles first

![Explore and place](v21_d_place.png){width=100%}

## Place tabs: Facts and Compare

![Facts tab](v21_d_facts.png){width=49%} ![Compare tab](v21_d_ptcompare.png){width=49%}

## Filters, records and the tools drawer

![Filters](v21_d_filters.png){width=100%}

![Records](v21_d_records.png){width=49%} ![Tools](v21_d_tools.png){width=49%}

## Compare three places

![Compare](v21_d_compare.png){width=100%}

## Higher or lower

![Game](v21_d_game.png){width=100%}

## Map layer: terrain

![Terrain layer](v21_d_layer.png){width=100%}

## Phone

![Home](v21_m_home.png){width=24%} ![Ranking](v21_m_list.png){width=24%} ![Place](v21_m_place.png){width=24%} ![Filters](v21_m_filters.png){width=24%}

![Tools](v21_m_tools.png){width=24%} ![Records](v21_m_records.png){width=24%} ![Map](v21_m_map.png){width=24%} ![Old v86](old_mobile_top.png){width=24%}
'''
if '2.2' in SRC:
    shots_md = """## Start screen: Germany and Pensioner couple chosen

![Start](v22_d_home.png){width=100%}

## After "Show my places": Germany opens

![Germany](v22_d_germany.png){width=100%}

## Greater Hamburg with a place open

![Greater Hamburg](v22_d_hamburg.png){width=100%}

## Ticking Compare, then the shortlist

![Ticked](v22_d_ticked.png){width=100%}

![Shortlist](v22_d_shortlist.png){width=100%}

## Compare three places

![Compare](v22_d_compare.png){width=100%}

## Everyday life within walking distance, and the wider region (simulated answers)

![Around](v22_d_around.png){width=49%} ![Facts](v22_d_facts.png){width=49%}

## Search: "Lisbon" finds the whole city

![Search](v22_d_search.png){width=100%}

## A greater-area page for search engines, and the terms of use

![Greater Hamburg page](v22_d_metropage.png){width=49%} ![Legal](v22_d_legal.png){width=49%}

## Phone

![Home](v22_m_home.png){width=24%} ![Map](v22_m_map.png){width=24%} ![Ranking](v22_m_list.png){width=24%} ![Place](v22_m_place.png){width=24%}

![Ticked](v22_m_ticked.png){width=24%} ![Around](v22_m_around.png){width=24%} ![Shortlist](v22_m_shortlist.png){width=24%} ![Filters](v22_m_filters.png){width=24%}

## The single file: Germany chosen on the start screen opens Germany

![Single file](sa_d_start_de.png){width=49%} ![Single file phone](sa_m_start_de.png){width=24%}
"""
if VER == '2.3':
    shots_md = """## New logo and the start screen

![Start](v23_d_home.png){width=100%}

## Explore: two whole top picks, the selected row in navy, factor detail only where you look

![Explore](v23_d_explore.png){width=100%}

## A place, and the profiles overview page

![Place](v23_d_place.png){width=100%}

![Profiles](v23_d_best.png){width=100%}

## Phone: profile pill, grouped picker, filter sheet on top, place

![Ranking](v23_m_list.png){width=24%} ![Picker](v23_m_picker.png){width=24%} ![Filters](v23_m_filters.png){width=24%} ![Place](v23_m_place.png){width=24%}

![Map](v23_m_map.png){width=24%} ![Scrolled: bars hide](v23_m_list_scrolled.png){width=24%}

## Phone held sideways

![Landscape](v23_l_list.png){width=100%}
"""
if VER == '2.4':
    shots_md = """## The privacy choice on the first visit (nothing is loaded from Google before it)

![Consent](v24_d_consent.png){width=100%}

## Country page: every place as a dot, the top 10 as linked pins

![Germany](v24_d_germany_map.png){width=100%}

![Norway, dark mode](v24_dk_norway_map.png){width=100%}

## Place page: strengths as bars, score bars, and the method page's data freshness

![Meilen](v24_d_meilen.png){width=100%}

![Freshness](v24_d_freshness.png){width=100%}

## The map with Tools open: Privacy settings and Imprint and privacy

![Tools](v24_d_tools.png){width=100%}

## Phone

![France map](v24_m_france_map.png){width=32%} ![Residents chart](v24_m_meilen_chart.png){width=32%} ![Ranking](v24_m_list.png){width=32%}
"""
rep = md.replace('SCREENSHOTS', shots_md)
open(os.path.join(ROOT, 'docs', SRC.replace('.md', '.filled.md')), 'w').write(md.replace('SCREENSHOTS', 'Screenshots are in the Word and PDF versions.'))
open(os.path.join(W, 'report.md'), 'w').write(rep)
dep = open(os.path.join(ROOT, 'docs', 'DEPLOY.md')).read()
dep = re.sub(r'^# .*\n', '', dep, count=1)
open(os.path.join(W, 'deploy.md'), 'w').write('---\ntitle: "Nestlyzer ' + VER + ': deploy guide"\nsubtitle: "Upload the sources once; GitHub builds, tests, publishes and refreshes the data"\ndate: "11 October 2026"\n---\n\n' + dep)

ref = os.path.join(W, 'ref.docx')
subprocess.run(['pandoc', '-o', ref, '--print-default-data-file', 'reference.docx'], check=True)
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document(ref)
for st in doc.styles:
    try:
        if st.font is not None: st.font.name = 'Arial'
    except Exception: pass
for name, size, col in [('Title', 24, '0B2E5C'), ('Subtitle', 13, '4A5A6E'), ('Heading 1', 17, '0B2E5C'), ('Heading 2', 13.5, '0B2E5C'), ('Heading 3', 11.5, '0B2E5C'), ('Body Text', 10.5, None), ('First Paragraph', 10.5, None), ('Compact', 10, None)]:
    try:
        s = doc.styles[name]; s.font.size = Pt(size)
        if col: s.font.color.rgb = RGBColor.from_string(col)
    except KeyError: pass
try:
    sc = doc.styles['Source Code']; sc.font.name = 'Courier New'; sc.font.size = Pt(8.5)
    pf = sc.element.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'EEF2F6'); pf.append(shd)
except KeyError: pass
for cs in ('Verbatim Char',):
    try: doc.styles[cs].font.name = 'Courier New'
    except KeyError: pass
doc.save(ref)
out = {'report': 'Nestlyzer-' + VER + '-Release-Report', 'deploy': 'Nestlyzer-' + VER + '-Deploy-Guide'}
for n, title in out.items():
    dx = os.path.join(W, n + '.docx')
    subprocess.run(['pandoc', n + '.md', '-o', dx, '--reference-doc=ref.docx', '--resource-path=.'], cwd=W, check=True)
    doc = Document(dx)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
    for t in doc.tables:
        bd = OxmlElement('w:tblBorders')
        for e in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            el = OxmlElement('w:' + e); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'C9D2DC'); bd.append(el)
        t._tbl.tblPr.append(bd)
        for i, row in enumerate(t.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5 if len(t.columns) >= 4 else 9.5)
                        if i == 0: r.font.bold = True
                if i == 0:
                    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'E8EEF4'); c._tc.get_or_add_tcPr().append(sh)
    doc.save(dx)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', W, dx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for ext in ('docx', 'pdf'):
        shutil.copy(os.path.join(W, n + '.' + ext), os.path.join(ROOT, 'dist', title + '.' + ext))
    print(title, 'ok')
